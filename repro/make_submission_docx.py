# -*- coding: utf-8 -*-
"""
make_submission_docx.py — converts the v0.5 manuscript and SI (markdown) into
Elsevier-style Word documents for submission (double/1.5-spaced, real tables).

Run with any python that has python-docx, e.g. the system Python 3.7:
  "C:\\Users\\15446\\AppData\\Local\\Programs\\Python\\Python37\\python.exe" repro/make_submission_docx.py
"""

import os
import re

from docx import Document
from docx.shared import Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))  # repo root
MAIN_MD = os.path.join(ROOT, "manuscript", "DILI_LABEL_STABILITY_DATA_NOTE_v0.5.md")
SI_MD = os.path.join(ROOT, "manuscript", "DILI_LABEL_STABILITY_DATA_NOTE_v0.5_SI.md")
OUT_DIR = os.path.join(ROOT, "submission_docx")
os.makedirs(OUT_DIR, exist_ok=True)

TOKEN = re.compile(r"(\*\*.+?\*\*|\*[^*\n]+?\*)")


def add_runs(par, text):
    for tok in TOKEN.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            r = par.add_run(tok[2:-2])
            r.bold = True
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            r = par.add_run(tok[1:-1])
            r.italic = True
        else:
            par.add_run(tok)


def add_para(doc, text, line_spacing=1.5):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_after = Pt(6)
    add_runs(p, text)
    return p


def is_sep(line):
    return bool(re.match(r"^\|[\s:\-|]+\|\s*$", line.strip()))


def render_table(doc, lines):
    rows = []
    for ln in lines:
        cells = [c.strip() for c in ln.strip().strip("|").split("|")]
        rows.append(cells)
    ncols = len(rows[0])
    table = doc.add_table(rows=0, cols=ncols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        tr = table.add_row()
        for j in range(ncols):
            celltxt = row[j] if j < len(row) else ""
            bold = "**" in celltxt or i == 0
            celltxt = celltxt.replace("**", "")
            tr.cells[j].text = celltxt
            for par in tr.cells[j].paragraphs:
                for run in par.runs:
                    if bold:
                        run.bold = True
    doc.add_paragraph()


def convert(src, out, title_prefix="", author_block=None):
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    lines = open(src, encoding="utf-8").read().splitlines()
    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        s = ln.strip()
        if not s or s == "---":
            i += 1
            continue
        if s.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not is_sep(lines[i]):
                    block.append(lines[i])
                i += 1
            render_table(doc, block)
            continue
        if s.startswith("# "):
            h = doc.add_heading(s[2:].strip(), level=0)
            for r in h.runs:
                r.font.size = Pt(16)
            i += 1
            continue
        if s.startswith("## "):
            doc.add_heading(s[3:].strip(), level=1)
            i += 1
            continue
        if s.startswith("### "):
            doc.add_heading(s[4:].strip(), level=2)
            i += 1
            continue
        if s.startswith("*Draft ") or s.startswith("*Remaining "):
            p = doc.add_paragraph()
            p.paragraph_format.line_spacing = 1.15
            add_runs(p, s)
            i += 1
            continue
        add_para(doc, s)
        i += 1
    doc.save(out)
    print("wrote", out)


convert(MAIN_MD, os.path.join(OUT_DIR, "manuscript.docx"))
convert(SI_MD, os.path.join(OUT_DIR, "supplementary_information.docx"))

# read-back verification of tables
doc = Document(os.path.join(OUT_DIR, "manuscript.docx"))
print("tables in manuscript.docx:", len(doc.tables))
for ti, t in enumerate(doc.tables):
    first = [c.text for c in t.rows[0].cells]
    print(" table", ti, "rows:", len(t.rows), "| header:", first)
si = Document(os.path.join(OUT_DIR, "supplementary_information.docx"))
print("tables in SI docx:", len(si.tables))
